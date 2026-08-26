from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import List

from ...base import BaseExportService
from ....models.AdvancesModel import AdvancesSupervisionExportRow

COLORS = {
    "header_bg": "1F3864",
    "header_fg": "FFFFFF",
    "row_even":  "F5F5F5",
    "row_odd":   "FFFFFF",
}

HEADERS = [
    ("N° Documento",    16),
    ("Tipo",            14),
    ("Empleado",        30),
    ("Cédula",          14),
    ("Valor solicitado", 16),
    ("Valor aprobado",   16),
    ("Estado",          22),
    ("Fecha solicitud",  16),
    ("Fecha apertura",   16),
    ("N° SAP",          16),
    ("Fecha legalización", 18),
]

DOCUMENT_TYPE_LABEL = {"ANTICIPO": "Anticipo", "CAJA_MENOR": "Caja Menor"}

STATUS_LABEL = {
    "PENDIENTE_NIVEL1": "Pendiente Nivel 1",
    "PENDIENTE_NIVEL2": "Pendiente Nivel 2 (Gerencia)",
    "PENDIENTE_VERIFICACION_CONTABLE": "Pendiente verificación contable",
    "PENDIENTE_TESORERIA": "Pendiente Tesorería",
    "ABIERTO": "Abierto",
    "PENDIENTE_APROBACION_LEGALIZACION": "Pendiente aprobación de legalización",
    "LEGALIZADO": "Legalizado",
    "RECHAZADO": "Rechazado",
}

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FONT = Font(bold=True, size=10, color=COLORS["header_fg"])
HEADER_FILL = PatternFill(start_color=COLORS["header_bg"], fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(size=9)
ALIGN_CENTER = Alignment(vertical="center", horizontal="center")
ALIGN_LEFT = Alignment(vertical="center", horizontal="left")
ROW_FILL_EVEN = PatternFill(start_color=COLORS["row_even"], fill_type="solid")
ROW_FILL_ODD = PatternFill(start_color=COLORS["row_odd"], fill_type="solid")

CENTERED_COLUMNS = {1, 2, 4, 5, 6, 7, 8, 9, 10, 11}


class AdvancesSupervisionExportService(BaseExportService):

    def generate_file(self, data: List[AdvancesSupervisionExportRow], options=None) -> io.BytesIO:
        wb = Workbook()
        self._build_sheet(wb, data, (options or {}).get("scope", ""))

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _build_sheet(self, wb: Workbook, rows: List[AdvancesSupervisionExportRow], scope: str):
        ws = wb.active
        ws.title = "Supervisión"
        ws.freeze_panes = "A3"

        title = "SUPERVISIÓN ANTICIPOS Y CAJA MENOR" if scope != "treasury" else "SUPERVISIÓN CAJA MENOR — TESORERÍA"
        ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
        title_cell = ws["A1"]
        title_cell.value = f"{title} — Erazo Valencia"
        title_cell.font = Font(bold=True, size=13, color=COLORS["header_fg"])
        title_cell.fill = HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        for col, (label, width) in enumerate(HEADERS, start=1):
            cell = ws.cell(row=2, column=col, value=label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[2].height = 24

        for i, row in enumerate(rows):
            r = i + 3
            fill = ROW_FILL_EVEN if i % 2 == 0 else ROW_FILL_ODD

            values = [
                row.documentNumber,
                DOCUMENT_TYPE_LABEL.get(row.documentType, row.documentType),
                row.employeeName,
                row.employeeDocumentId,
                row.requestedAmount,
                row.approvedAmount if row.approvedAmount is not None else "",
                STATUS_LABEL.get(row.status, row.status),
                row.createdAt,
                row.openedAt,
                row.erpAdvanceNumber,
                row.legalizationApprovedAt,
            ]

            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.fill = fill
                cell.border = BORDER
                cell.alignment = ALIGN_CENTER if col in CENTERED_COLUMNS else ALIGN_LEFT
                cell.font = DATA_FONT
                if col in (5, 6):
                    cell.number_format = "#,##0"

            ws.row_dimensions[r].height = 16

        ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{len(rows) + 2}"

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
