from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import List

from ...base import BaseExportService
from ....models.AdvancesModel import AdvancesInvoiceExportRow

COLORS = {
    "header_bg": "1F3864",
    "header_fg": "FFFFFF",
    "row_even":  "F5F5F5",
    "row_odd":   "FFFFFF",
}

HEADERS = [
    ("N° Factura",     16),
    ("Se carga a",     16),
    ("Tipo",           16),
    ("NIT/CC Tercero", 18),
    ("Nombre Tercero", 30),
    ("Valor",          14),
    ("Concepto",       36),
    ("Legalización",   16),
    ("Contable",       16),
    ("Registrada por", 26),
    ("Fecha",          14),
]

CHARGE_NATURE_LABEL = {"ANTICIPO": "Anticipo", "CAJA_MENOR": "Caja Menor", "TC": "Tarjeta de Crédito", "NINGUNA": "Sin asociar"}
STATUS1_LABEL = {"POR_LEGALIZAR": "Por legalizar", "LEGALIZADO": "Legalizado"}
STATUS2_LABEL = {"POR_VERIFICAR": "Por verificar", "VERIFICADA": "Verificada", "CONTABILIZADA": "Contabilizada", "RECHAZADA": "Rechazada"}

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FONT = Font(bold=True, size=10, color=COLORS["header_fg"])
HEADER_FILL = PatternFill(start_color=COLORS["header_bg"], fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(size=9)
ALIGN_CENTER = Alignment(vertical="center", horizontal="center")
ALIGN_LEFT = Alignment(vertical="center", horizontal="left")
ROW_FILL_EVEN = PatternFill(start_color=COLORS["row_even"], fill_type="solid")
ROW_FILL_ODD = PatternFill(start_color=COLORS["row_odd"], fill_type="solid")

CENTERED_COLUMNS = {1, 2, 3, 4, 6, 8, 9, 11}


class AdvancesInvoiceExportService(BaseExportService):

    def generate_file(self, data: List[AdvancesInvoiceExportRow], options=None) -> io.BytesIO:
        wb = Workbook()
        self._build_sheet(wb, data)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _build_sheet(self, wb: Workbook, rows: List[AdvancesInvoiceExportRow]):
        ws = wb.active
        ws.title = "Facturas"
        ws.freeze_panes = "A3"

        ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
        title_cell = ws["A1"]
        title_cell.value = "FACTURAS DE LEGALIZACIÓN — Anticipos/Caja Menor/TC — Erazo Valencia"
        title_cell.font = Font(bold=True, size=13, color=COLORS["header_fg"])
        title_cell.fill = HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        for col, (label, width) in enumerate(HEADERS, start=1):
            cell = ws.cell(row=2, column=col, value=label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[2].height = 24

        for i, row in enumerate(rows):
            r = i + 3
            fill = ROW_FILL_EVEN if i % 2 == 0 else ROW_FILL_ODD

            values = [
                row.invoiceNumber,
                row.documentNumber,
                CHARGE_NATURE_LABEL.get(row.chargeNature, row.chargeNature),
                row.thirdPartyTaxId,
                row.thirdPartyName,
                row.amount,
                row.concept,
                STATUS1_LABEL.get(row.status1, row.status1),
                STATUS2_LABEL.get(row.status2, row.status2),
                row.registeredByName,
                row.createdAt,
            ]

            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.fill = fill
                cell.border = BORDER
                cell.alignment = ALIGN_CENTER if col in CENTERED_COLUMNS else ALIGN_LEFT
                cell.font = DATA_FONT
                if col == 6:
                    cell.number_format = "#,##0"

            ws.row_dimensions[r].height = 16

        ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{len(rows) + 2}"

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
