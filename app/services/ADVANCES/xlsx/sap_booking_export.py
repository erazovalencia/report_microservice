from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import io
from typing import List

from ...base import BaseExportService
from ....models.AdvancesModel import AdvancesSapBookingRow

# Layout real de 43 columnas de la plantilla de cargue FI que entregó
# Gerencia (MASIVOfinal2.xlsx, hoja "Plantilla" — ver
# ANTICIPOS_SAP_CARGUE_DESIGN.md §4). 3 filas de metadatos antes de los
# datos reales, igual que el archivo original:
#   fila 1 = nombre humano del campo
#   fila 2 = código técnico SAP (lo que SAP realmente lee)
#   fila 3 = longitud/tipo de campo SAP — restricción dura, no negociable
# Solo el bloque 1 (columnas P-V) está en uso hoy; los bloques 2-4
# (columnas W-AQ) quedan reservados y siempre en blanco (§5.1).
SAP_COLUMNS = [
    ("FI: Relación de operaciones contables", "BUSCS", "C(001)"),
    ("Cuenta o matchcode para la siguiente posición", "ACCNT", "C(010)"),
    ("Fe.factura en documento", "BLDAT", "C(010)"),
    ("Número de documento de referencia", "XBLNR", "C(016)"),
    ("Fecha de contabilización en el documento", "BUDAT", "C(010)"),
    ("Clase de documento", "BLART", "C(002)"),
    ("Importe en la moneda del documento", "WRBTR", "C(013)"),
    ("Clave de moneda", "WAERS", "C(003)"),
    ("¿Calcular impuesto automáticamente?", "XMWST", "C(001)"),
    ("Indicador IVA", "MWSKZ", "C(002)"),
    ("Texto posición", "SGTXT", "C(025)"),
    ("Base imponible de retención en moneda de documento", "WT_QSSHB_01", "C(013)"),
    ("Base imponible de retención en moneda de documento", "WT_QSSHB_02", "C(013)"),
    ("Importe de retención en moneda de documento", "WT_QBSHB_01", "C(013)"),
    ("Importe de retención en moneda de documento", "WT_QBSHB_02", "C(013)"),
    ("Cuenta de mayor de la contabilidad principal 1", "HKONT_01", "C(010)"),
    ("Importe en la moneda del documento 1", "WRBTR_01", "C(013)"),
    ("Indicador IVA 1", "MWSKZ_01", "C(002)"),
    ("Texto posición 1", "SGTXT_01", "C(050)"),
    ("Centro de coste 1", "KOSTL_01", "C(010)"),
    ("Número de orden 1", "AUFNR_01", "C(012)"),
    ("Elemento del plan de estructura de proyecto (elemento PEP) 1", "PROJK_01", "C(024)"),
    ("Cuenta de mayor de la contabilidad principal 2", "HKONT_02", "C(010)"),
    ("Importe en la moneda del documento 2", "WRBTR_02", "C(013)"),
    ("Indicador IVA 2", "MWSKZ_02", "C(002)"),
    ("Texto posición 2", "SGTXT_02", "C(050)"),
    ("Centro de coste 2", "KOSTL_02", "C(010)"),
    ("Número de orden 2", "AUFNR_02", "C(012)"),
    ("Elemento del plan de estructura de proyecto (elemento PEP) 2", "PROJK_02", "C(024)"),
    ("Cuenta de mayor de la contabilidad principal 3", "HKONT_03", "C(010)"),
    ("Importe en la moneda del documento 3", "WRBTR_03", "C(013)"),
    ("Indicador IVA 3", "MWSKZ_03", "C(002)"),
    ("Texto posición 3", "SGTXT_03", "C(050)"),
    ("Centro de coste 3", "KOSTL_03", "C(010)"),
    ("Número de orden 3", "AUFNR_03", "C(012)"),
    ("Elemento del plan de estructura de proyecto (elemento PEP) 3", "PROJK_03", "C(024)"),
    ("Cuenta de mayor de la contabilidad principal 4", "HKONT_04", "C(010)"),
    ("Importe en la moneda del documento 4", "WRBTR_04", "C(013)"),
    ("Indicador IVA 4", "MWSKZ_04", "C(002)"),
    ("Texto posición 4", "SGTXT_04", "C(050)"),
    ("Centro de coste 4", "KOSTL_04", "C(010)"),
    ("Número de orden 4", "AUFNR_04", "C(012)"),
    ("Elemento del plan de estructura de proyecto (elemento PEP) 4", "PROJK_04", "C(024)"),
]
TOTAL_COLUMNS = len(SAP_COLUMNS)  # 43 (A..AQ)

# Estilos reutilizados por referencia — instanciar uno nuevo por celda/fila es
# el patrón más lento conocido en openpyxl (decisión 2026-07-27,
# RdpReportExportService).
HEADER_NAME_FONT = Font(bold=True, size=9)
HEADER_CODE_FONT = Font(bold=True, size=9, color="1F3864")
HEADER_FORMAT_FONT = Font(italic=True, size=8, color="808080")
DATA_FONT = Font(size=9)
ALIGN_WRAP = Alignment(vertical="center", horizontal="left", wrap_text=True)
ALIGN_LEFT = Alignment(vertical="center", horizontal="left")


class AdvancesSapBookingExportService(BaseExportService):
    """
    Genera el archivo de cargue masivo a SAP en el layout real de 43
    columnas (ANTICIPOS_SAP_CARGUE_DESIGN.md, Etapa 3). Cada fila de entrada
    ya trae los valores resueltos por valera (match de acreedor, cuenta
    contable del servicio, centro de costo, texto armado) — este servicio
    solo vuelca esos valores en las columnas A-V exactas; las columnas
    W-AQ (bloques 2-4) quedan vacías, reservadas para facturas multi-línea
    contable fuera de alcance de esta entrega.
    """

    def generate_file(self, data: List[AdvancesSapBookingRow], options=None) -> io.BytesIO:
        wb = Workbook()
        self._build_sheet(wb, data)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _build_sheet(self, wb: Workbook, rows: List[AdvancesSapBookingRow]):
        ws = wb.active
        ws.title = "Cargue SAP"
        ws.freeze_panes = "A4"

        for col, (human_name, sap_code, fmt) in enumerate(SAP_COLUMNS, start=1):
            letter = get_column_letter(col)

            name_cell = ws.cell(row=1, column=col, value=human_name)
            name_cell.font = HEADER_NAME_FONT
            name_cell.alignment = ALIGN_WRAP

            code_cell = ws.cell(row=2, column=col, value=sap_code)
            code_cell.font = HEADER_CODE_FONT
            code_cell.alignment = ALIGN_LEFT

            fmt_cell = ws.cell(row=3, column=col, value=fmt)
            fmt_cell.font = HEADER_FORMAT_FONT
            fmt_cell.alignment = ALIGN_LEFT

            # Columnas en uso (A-V) más anchas para lectura; el resto
            # (bloques 2-4, siempre vacíos) angostas para no ocupar pantalla.
            ws.column_dimensions[letter].width = 20 if col <= 22 else 10

        ws.row_dimensions[1].height = 30

        for i, row in enumerate(rows):
            r = i + 4  # los datos empiezan en la fila 4, igual que el archivo real de Gerencia
            values = [
                row.buscs,
                row.accnt,
                row.bldat,
                row.xblnr,
                row.budat,
                row.blart,
                row.wrbtr,
                row.waers,
                row.xmwst,
                row.mwskz,
                row.sgtxt,
                None, None, None, None,  # L, M, N, O — retención, sin uso en v1 (§4)
                row.hkont_01,
                row.wrbtr_01,
                row.mwskz_01,
                row.sgtxt_01,
                row.kostl_01,
                row.aufnr_01,
                row.projk_01,
                # W..AQ (bloques 2-4) — reservados, siempre vacíos (§5.1)
            ]

            for col, val in enumerate(values, start=1):
                if val is None or val == "":
                    continue
                cell = ws.cell(row=r, column=col, value=val)
                cell.font = DATA_FONT
                cell.alignment = ALIGN_LEFT
                if col in (7, 17):  # WRBTR, WRBTR_01 — valores monetarios
                    cell.number_format = "#,##0.00"

            ws.row_dimensions[r].height = 15

        last_row = len(rows) + 3
        if last_row >= 3:
            ws.auto_filter.ref = f"A3:{get_column_letter(TOTAL_COLUMNS)}{last_row}"

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
