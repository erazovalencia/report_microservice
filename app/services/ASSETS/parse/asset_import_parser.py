import csv
import io
import re
from datetime import datetime
from typing import List, Dict, Any, Optional, Iterable, Sequence, Tuple
import openpyxl

from ._text import norm as _norm, norm_header as _norm_header, headers as _headers, find_col as _find_col

# Import masivo de activos desde un export de equipos de SAP (ver
# valera/ASSETS_INVENTORY_DESIGN.md §2/§6). Acepta dos formatos: el .xlsx de
# IH08 y el export de lista de SAP guardado como ".XLS" (equipos2309.XLS), que
# en realidad es texto UTF-16 separado por tabuladores — el contenido se
# detecta por sus bytes, nunca por la extensión. Detección de columnas por
# encabezado, no por nombre/posición de hoja (mismo criterio que
# balance_import_parser.py); ambos formatos usan encabezados distintos para el
# mismo dato ("Equipo superior" / "EquiSuper."), por eso cada campo admite
# varios alias.


EQUIPMENT_HEADERS = _headers("Equipo")
PARENT_EQUIPMENT_HEADERS = _headers("Equipo superior", "EquiSuper.")
FIXED_ASSET_HEADERS = _headers("Activo fijo")
MANUFACTURER_HEADERS = _headers("Fabricante")
# MAESTRO AF toma SERIAL de "Fabr. Nº-serie" (serial del fabricante) y deja
# "Número de serie" (identificador propio de SAP) como campo separado.
MANUFACTURER_SERIAL_HEADERS = _headers("Fabr. Nº-serie", "Nº serie fabricante")
PLAIN_SERIAL_HEADERS = _headers("Número de serie", "NºSerie")
COMPANY_HEADERS = _headers("Sociedad", "Soc.")
COST_CENTER_HEADERS = _headers("Centro coste", "Centro de costo", "Ce.coste")
PEP_HEADERS = _headers("Elemento PEP")
# "Denominación" aparece 2 veces en IH08: la del Equipo (justo después de
# "Fabr. Nº-serie") y la del Material (justo después de la columna "Material").
# _find_col toma la primera coincidencia = la del equipo.
DESCRIPTION_HEADERS = _headers("Denominación", "Denominación de objeto técnico")
MATERIAL_DESCRIPTION_HEADERS = DESCRIPTION_HEADERS | _headers("Texto breve de material")
PART_NUMBER_HEADERS = _headers(
    "NºPieza fabric.", "Nº Pieza fabric.", "Número de pieza de fabricante"
)
INVENTORY_NUMBER_HEADERS = _headers("Nº inventario")
LOCATION_HEADERS = _headers("Local")
TECHNICAL_LOCATION_HEADERS = _headers("Ubicac.técnica", "Ubicación técnica")
SUB_NUMBER_HEADERS = _headers("Subnúmero", "SNº")
EMPLACEMENT_HEADERS = _headers("Emplazamiento", "Emplaz.")
# "Ce." del export de lista = centro del emplazamiento (agrupador de emplazamientos).
EMPLACEMENT_CENTER_HEADERS = _headers("Ce.emplazam.", "Centro del emplazamiento", "Ce.")
MATERIAL_HEADERS = _headers("Material")
MODIFIED_AT_HEADERS = _headers("Modificado el", "Modif.el")
MODIFIED_BY_HEADERS = _headers("Modificado por", "Modif.por")
SYSTEM_STATUS_HEADERS = _headers("Status sistema", "Stat.sist.")

# "1,52005E+11": Excel/SAP convirtió el Activo fijo a notación científica y la
# precisión ya se perdió en el export (22 valores distintos en 8.656 filas) — no
# es recuperable, mejor vacío que un dato falso. Hay que re-exportar la columna
# como texto si se necesita.
_SCIENTIFIC_NOTATION = re.compile(r"\d+[.,]\d+[eE][+-]?\d+")


def _norm_code(val) -> str:
    """SAP entrega 'Equipo'/'Activo fijo' como número — sin '.0' final."""
    s = _norm(val)
    if not s:
        return ""
    try:
        return str(int(float(s)))
    except (ValueError, TypeError):
        return s


def _norm_fixed_asset(val) -> str:
    s = _norm_code(val)
    return "" if _SCIENTIFIC_NOTATION.fullmatch(s) else s


def _norm_date(val) -> Optional[str]:
    """openpyxl (data_only=True) entrega celdas de fecha como datetime — ISO
    date (sin hora, IH08 no la trae con significado) para que Node la parsee
    con `new Date()` sin depender de la zona horaria de quien corre el import.
    El export de texto trae DD.MM.AAAA."""
    if val is None or val == "":
        return None
    if hasattr(val, "date"):
        return val.date().isoformat()
    if hasattr(val, "isoformat"):
        return val.isoformat()
    s = _norm(val)
    match = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", s)
    if match:
        try:
            return datetime(int(match[3]), int(match[2]), int(match[1])).date().isoformat()
        except ValueError:
            return None
    return s or None


def _locate_columns(headers: List[str]) -> Optional[Dict[str, Any]]:
    """headers ya normalizados con _norm_header. None si no hay columna 'Equipo'."""
    equipment_idx = _find_col(headers, EQUIPMENT_HEADERS)
    if equipment_idx is None:
        return None

    manufacturer_serial_idx = _find_col(headers, MANUFACTURER_SERIAL_HEADERS)
    plain_serial_idx = _find_col(headers, PLAIN_SERIAL_HEADERS)
    material_idx = _find_col(headers, MATERIAL_HEADERS)
    # La descripción del Material es la columna siguiente a "Material" en el
    # layout real — mismo texto de encabezado que la del equipo en IH08, se
    # distingue solo por posición.
    material_description_idx = (
        material_idx + 1
        if material_idx is not None
        and material_idx + 1 < len(headers)
        and headers[material_idx + 1] in MATERIAL_DESCRIPTION_HEADERS
        else None
    )

    description_idx = _find_col(headers, DESCRIPTION_HEADERS)
    # El export de texto reserva columnas SIN encabezado justo después de la
    # denominación; si el texto trae tabuladores, el resto cae ahí (ej.
    # 'LUBRICATOR' | '7" X 120' | 'INTEGRAL' | '10000'). Se re-unen al texto.
    description_continuation: List[int] = []
    if description_idx is not None:
        j = description_idx + 1
        while j < len(headers) and headers[j] == "":
            description_continuation.append(j)
            j += 1

    return {
        "equipment": equipment_idx,
        "parentEquipment": _find_col(headers, PARENT_EQUIPMENT_HEADERS),
        "fixedAsset": _find_col(headers, FIXED_ASSET_HEADERS),
        "manufacturer": _find_col(headers, MANUFACTURER_HEADERS),
        "serial": manufacturer_serial_idx if manufacturer_serial_idx is not None else plain_serial_idx,
        "company": _find_col(headers, COMPANY_HEADERS),
        "costCenter": _find_col(headers, COST_CENTER_HEADERS),
        "pep": _find_col(headers, PEP_HEADERS),
        "description": description_idx,
        "descriptionContinuation": description_continuation,
        "partNumber": _find_col(headers, PART_NUMBER_HEADERS),
        "plainSerial": plain_serial_idx,
        "inventoryNumber": _find_col(headers, INVENTORY_NUMBER_HEADERS),
        "location": _find_col(headers, LOCATION_HEADERS),
        "technicalLocation": _find_col(headers, TECHNICAL_LOCATION_HEADERS),
        "subNumber": _find_col(headers, SUB_NUMBER_HEADERS),
        "emplacement": _find_col(headers, EMPLACEMENT_HEADERS),
        "emplacementCenter": _find_col(headers, EMPLACEMENT_CENTER_HEADERS),
        "material": material_idx,
        "materialDescription": material_description_idx,
        "modifiedAt": _find_col(headers, MODIFIED_AT_HEADERS),
        "modifiedBy": _find_col(headers, MODIFIED_BY_HEADERS),
        "systemStatus": _find_col(headers, SYSTEM_STATUS_HEADERS),
    }


Table = Tuple[Dict[str, Any], Iterable[Tuple[int, Sequence]]]


def _table_from_workbook(file_bytes: bytes) -> Optional[Table]:
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
    for ws in wb.worksheets:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            continue
        cols = _locate_columns([_norm_header(c) for c in first_row])
        if cols is None:
            continue
        equipment_idx = cols["equipment"]
        has_data = any(
            len(row) > equipment_idx and _norm(row[equipment_idx]) != ""
            for row in ws.iter_rows(min_row=2, values_only=True)
        )
        if has_data:
            return cols, enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
    return None


def _decode_text_export(file_bytes: bytes) -> str:
    if file_bytes[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return file_bytes.decode("utf-16")
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return file_bytes.decode("cp1252")


def _table_from_text_export(file_bytes: bytes) -> Optional[Table]:
    # Dialecto CSV estándar: SAP entrecomilla (con "" para comillas internas) los
    # campos que traen tabuladores o comillas, ej. '" 7"" X 120"'. Una comilla
    # en medio de un campo sin entrecomillar (7" X 120) se lee como literal.
    reader = csv.reader(io.StringIO(_decode_text_export(file_bytes), newline=""), delimiter="\t")
    lines = list(reader)
    if not lines:
        return None
    cols = _locate_columns([_norm_header(c) for c in lines[0]])
    if cols is None:
        return None
    return cols, enumerate(lines[1:], start=2)


def _is_binary_workbook(file_bytes: bytes) -> bool:
    # .xlsx/.xlsm = zip ("PK"); .xls binario = OLE2. Cualquier otra cosa es el
    # export de lista de SAP (texto), aunque venga con extensión .XLS.
    return file_bytes[:2] == b"PK" or file_bytes[:4] == b"\xd0\xcf\x11\xe0"


def _cell(row: Sequence, idx: Optional[int]):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def parse_asset_import_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    table = (
        _table_from_workbook(file_bytes)
        if _is_binary_workbook(file_bytes)
        else _table_from_text_export(file_bytes)
    )
    if table is None:
        raise ValueError(
            "No se encontró una columna 'Equipo' (export de equipos de SAP: IH08 .xlsx o lista .XLS)"
        )
    cols, numbered_rows = table

    rows_out: List[Dict[str, Any]] = []
    for row_number, row in numbered_rows:
        if all(v is None or _norm(v) == "" for v in row):
            continue

        errors: List[str] = []

        sap_equipment_code = _norm_code(_cell(row, cols["equipment"]))
        if not sap_equipment_code:
            errors.append("Equipo (código SAP) vacío")

        parent_sap_equipment_code = _norm_code(_cell(row, cols["parentEquipment"])) or None
        sap_fixed_asset_number = _norm_fixed_asset(_cell(row, cols["fixedAsset"])) or None
        brand = _norm(_cell(row, cols["manufacturer"])) or None
        serial_number = _norm(_cell(row, cols["serial"])) or None
        sap_company = _norm(_cell(row, cols["company"])) or None
        sap_cost_center = _norm(_cell(row, cols["costCenter"])) or None
        sap_pep_element = _norm(_cell(row, cols["pep"])) or None

        description = " ".join(
            part
            for part in (
                _norm(_cell(row, cols["description"])),
                *(_norm(_cell(row, i)) for i in cols["descriptionContinuation"]),
            )
            if part
        ) or None
        part_number = _norm(_cell(row, cols["partNumber"])) or None
        internal_serial_number = _norm(_cell(row, cols["plainSerial"])) or None
        sap_inventory_number = _norm(_cell(row, cols["inventoryNumber"])) or None
        sap_location = _norm(_cell(row, cols["location"])) or None
        sap_technical_location = _norm(_cell(row, cols["technicalLocation"])) or None
        sub_number = _norm(_cell(row, cols["subNumber"])) or None
        emplacement_code = _norm(_cell(row, cols["emplacement"])) or None
        emplacement_center = _norm(_cell(row, cols["emplacementCenter"])) or None
        material_code = _norm_code(_cell(row, cols["material"])) or None
        material_description = _norm(_cell(row, cols["materialDescription"])) or None
        sap_modified_at = _norm_date(_cell(row, cols["modifiedAt"]))
        sap_modified_by = _norm(_cell(row, cols["modifiedBy"])) or None
        sap_system_status = _norm(_cell(row, cols["systemStatus"])) or None

        rows_out.append({
            "rowIndex": row_number,
            "sapEquipmentCode": sap_equipment_code,
            "parentSapEquipmentCode": parent_sap_equipment_code,
            "sapFixedAssetNumber": sap_fixed_asset_number,
            "brand": brand,
            "serialNumber": serial_number,
            "sapCostCenter": sap_cost_center,
            "sapPepElement": sap_pep_element,
            "description": description,
            "partNumber": part_number,
            "internalSerialNumber": internal_serial_number,
            "sapInventoryNumber": sap_inventory_number,
            "sapCompany": sap_company,
            "sapLocation": sap_location,
            "sapTechnicalLocation": sap_technical_location,
            "subNumber": sub_number,
            "emplacementCode": emplacement_code,
            "emplacementCenter": emplacement_center,
            "materialCode": material_code,
            "materialDescription": material_description,
            "sapModifiedAt": sap_modified_at,
            "sapModifiedBy": sap_modified_by,
            "sapSystemStatus": sap_system_status,
            "parseErrors": errors,
        })

    return rows_out
