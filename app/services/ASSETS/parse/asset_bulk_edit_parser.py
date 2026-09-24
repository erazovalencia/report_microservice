import io
from typing import List, Dict, Any, Optional
import openpyxl

from ._text import norm, norm_header, headers, find_col

# Reimport del propio "Exportar Excel" del panel (ver AssetExportService en
# report-microservice) para edición masiva — deliberadamente SEPARADO del
# import SAP (asset_import_parser.py): ese matchea por sapEquipmentCode y
# refresca los campos de origen SAP. Este matchea por assetTag ("Nº
# inventario", el ID interno de VALERA, único).
#
# Alcance acotado a propósito (2026-09-24, decisión explícita): Denominación,
# Local, Ubicación técnica, Emplazamiento, Ce.emplazam., Elemento PEP, Centro
# coste, Sociedad y Equipo superior — el resto de columnas del export (Línea,
# Categoría, Estado, Fabricante, Modelo, Serial, Nº inventario SAP, Status
# sistema SAP, Asignado a) quedan de solo lectura en este import hasta que se
# decida ampliar el alcance. "Equipo superior" viaja como código SAP del padre
# (igual que en IH08); el handler lo resuelve a un activo y valida ciclos.

HEADER_SCAN_ROWS = 10
ASSET_TAG_HEADERS = headers("Nº inventario")
DESCRIPTION_HEADERS = headers("Denominación")
LOCATION_HEADERS = headers("Local")
TECHNICAL_LOCATION_HEADERS = headers("Ubicación técnica", "Ubicac.técnica")
EMPLACEMENT_HEADERS = headers("Emplazamiento", "Emplaz.")
EMPLACEMENT_CENTER_HEADERS = headers("Ce.emplazam.", "Centro del emplazamiento")
PEP_HEADERS = headers("Elemento PEP")
COST_CENTER_HEADERS = headers("Centro coste", "Ce.coste")
COMPANY_HEADERS = headers("Sociedad", "Soc.")
PARENT_EQUIPMENT_HEADERS = headers("Equipo superior", "EquiSuper.")


def _locate_columns(row_headers: List[str]) -> Optional[Dict[str, Any]]:
    asset_tag_idx = find_col(row_headers, ASSET_TAG_HEADERS)
    if asset_tag_idx is None:
        return None
    return {
        "assetTag": asset_tag_idx,
        "description": find_col(row_headers, DESCRIPTION_HEADERS),
        "sapLocation": find_col(row_headers, LOCATION_HEADERS),
        "sapTechnicalLocation": find_col(row_headers, TECHNICAL_LOCATION_HEADERS),
        "emplacementCode": find_col(row_headers, EMPLACEMENT_HEADERS),
        "emplacementCenter": find_col(row_headers, EMPLACEMENT_CENTER_HEADERS),
        "sapPepElement": find_col(row_headers, PEP_HEADERS),
        "sapCostCenter": find_col(row_headers, COST_CENTER_HEADERS),
        "sapCompany": find_col(row_headers, COMPANY_HEADERS),
        "parentEquipment": find_col(row_headers, PARENT_EQUIPMENT_HEADERS),
    }


def _cell(row, idx: Optional[int]):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _code(val) -> Optional[str]:
    """Códigos de SAP (1000, 1000000) que Excel puede devolver como 1000.0 al
    editarlos como número — sin el '.0' final."""
    s = norm(val)
    if not s:
        return None
    try:
        return str(int(float(s))) if float(s).is_integer() else s
    except (ValueError, TypeError):
        return s


def parse_asset_bulk_edit_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)

    # El export lleva un banner de título en la fila 1 y los encabezados en la
    # fila 2 — se busca la fila de encabezados en las primeras filas, no se asume
    # la 1. Los números de fila que se reportan son los reales de la hoja.
    cols = None
    rows_iter = None
    for ws in wb.worksheets:
        for header_row_number, header_row in enumerate(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
            found = _locate_columns([norm_header(c) for c in header_row])
            if found is not None:
                cols = found
                rows_iter = enumerate(
                    ws.iter_rows(min_row=header_row_number + 1, values_only=True), start=header_row_number + 1
                )
                break
        if cols is not None:
            break

    if cols is None:
        raise ValueError(
            "No se encontró la columna 'Nº inventario' — sube el mismo archivo que descargaste con \"Exportar Excel\"."
        )

    rows_out: List[Dict[str, Any]] = []
    for row_number, row in rows_iter:
        if all(v is None or norm(v) == "" for v in row):
            continue

        errors: List[str] = []
        asset_tag = norm(_cell(row, cols["assetTag"]))
        if not asset_tag:
            errors.append("Nº inventario vacío")

        rows_out.append({
            "rowIndex": row_number,
            "assetTag": asset_tag,
            "description": norm(_cell(row, cols["description"])) or None,
            "sapLocation": norm(_cell(row, cols["sapLocation"])) or None,
            "sapTechnicalLocation": norm(_cell(row, cols["sapTechnicalLocation"])) or None,
            "emplacementCode": _code(_cell(row, cols["emplacementCode"])),
            "emplacementCenter": _code(_cell(row, cols["emplacementCenter"])),
            "sapPepElement": norm(_cell(row, cols["sapPepElement"])) or None,
            "sapCostCenter": _code(_cell(row, cols["sapCostCenter"])),
            "sapCompany": _code(_cell(row, cols["sapCompany"])),
            "parentEquipment": _code(_cell(row, cols["parentEquipment"])),
            "parseErrors": errors,
        })

    return rows_out
