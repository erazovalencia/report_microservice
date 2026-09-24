from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import List

from ...base import BaseExportService
from ....models.AssetModel import AssetExportRow

COLORS = {
    "header_bg": "1F3864",
    "header_fg": "FFFFFF",
    "row_even":  "F5F5F5",
    "row_odd":   "FFFFFF",
}

# Mismos labels en español que src/lib/assetStatusLabels.ts en valera — se
# duplican acá (no hay forma de compartir código TS/Python), mismo criterio
# que TIPO_LABEL/ESTADO_LABEL de RDP/Asistencia.
ASSET_LINE_LABEL = {
    "EMPLOYEE_DEVICE": "Dispositivo de empleado",
    "FIELD_EQUIPMENT": "Equipo pesado de campo",
}

ASSET_STATUS_LABEL = {
    "OPERATIVO": "Operativo",
    "EN_MANTENIMIENTO": "En mantenimiento",
    "EN_CONSTRUCCION": "En construcción",
    "INCOMPLETO": "Incompleto",
    "POR_UBICAR": "Por ubicar",
    "FUERA_DE_SERVICIO": "Fuera de servicio",
    "BAJA": "Baja",
}

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

HEADERS = [
    ("Nº inventario",      16),
    ("Equipo (SAP)",       12),
    ("Equipo superior",    14),
    ("Línea",              22),
    ("Categoría",          26),
    ("Denominación",       34),
    ("Fabricante",         18),
    ("Modelo",             18),
    ("Serial",             18),
    ("Nº inventario SAP",  16),
    ("Local",              14),
    ("Ubicación técnica",  22),
    ("Emplazamiento",      14),
    ("Ce.emplazam.",       12),
    ("Sociedad",           10),
    ("Centro coste",       12),
    ("Elemento PEP",       24),
    ("Estado",             16),
    ("Status sistema SAP", 16),
    ("Asignado a",         22),
    ("Creado",             12),
]

# Estilos de celda reutilizados por referencia — instanciar Font/Alignment por
# celda es el patrón más lento de openpyxl (ver fix de RdpReportExportService,
# 2026-07-27); con un catálogo de ~21k filas esto no es opcional acá.
DATA_FONT = Font(size=9)
ALIGN_CENTER = Alignment(vertical="center", horizontal="center")
ALIGN_LEFT = Alignment(vertical="center", horizontal="left")
CENTERED_COLS = {2, 3, 4, 13, 14, 15, 16, 18, 19, 21}
FILL_EVEN = PatternFill(start_color=COLORS["row_even"], fill_type="solid")
FILL_ODD = PatternFill(start_color=COLORS["row_odd"], fill_type="solid")


class AssetExportService(BaseExportService):

    def generate_file(self, data: List[AssetExportRow], options=None) -> io.BytesIO:
        wb = Workbook()
        self._build_sheet(wb, data)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _build_sheet(self, wb: Workbook, rows: List[AssetExportRow]):
        ws = wb.active
        ws.title = "Inventario de Activos"
        ws.freeze_panes = "A3"

        ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
        title_cell = ws["A1"]
        title_cell.value = "INVENTARIO DE ACTIVOS — Erazo Valencia"
        title_cell.font = Font(bold=True, size=13, color=COLORS["header_fg"])
        title_cell.fill = PatternFill(start_color=COLORS["header_bg"], fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        for col, (label, width) in enumerate(HEADERS, start=1):
            cell = ws.cell(row=2, column=col, value=label)
            cell.font = Font(bold=True, size=10, color=COLORS["header_fg"])
            cell.fill = PatternFill(start_color=COLORS["header_bg"], fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[2].height = 24

        for i, row in enumerate(rows):
            r = i + 3
            fill = FILL_EVEN if i % 2 == 0 else FILL_ODD

            values = [
                row.assetTag,
                row.sapEquipmentCode or "",
                row.parentEquipment or "",
                ASSET_LINE_LABEL.get(row.assetLine, row.assetLine),
                row.category or "",
                row.description or "",
                row.brand or "",
                row.model or "",
                row.serialNumber or "",
                row.sapInventoryNumber or "",
                row.sapLocation or "",
                row.sapTechnicalLocation or "",
                row.emplacementCode or "",
                row.emplacementCenter or "",
                row.sapCompany or "",
                row.sapCostCenter or "",
                row.sapPepElement or "",
                ASSET_STATUS_LABEL.get(row.status, row.status),
                row.sapSystemStatus or "",
                row.assignedTo or "",
                (row.createdAt or "")[:10],
            ]

            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.fill = fill
                cell.border = BORDER
                cell.font = DATA_FONT
                cell.alignment = ALIGN_CENTER if col in CENTERED_COLS else ALIGN_LEFT

            ws.row_dimensions[r].height = 16

        ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{len(rows) + 2}"

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
