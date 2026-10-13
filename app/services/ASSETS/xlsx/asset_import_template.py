from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import io

from ...base import BaseExportService
from ..parse.asset_import_parser import ASSET_IMPORT_TEMPLATE_SHEET

# Plantilla de alta de activos para "Importar masivo". Los encabezados son los
# mismos alias de IH08 que reconoce asset_import_parser.py, en la fila 1 y sin
# fila de ejemplo: el parser lee como dato todo lo que esté debajo del
# encabezado, así que las indicaciones van en comentarios de celda y en la hoja
# de instrucciones.

HEADER_BG = "1F3864"
HEADER_FG = "FFFFFF"
REQUIRED_BG = "C55A11"

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Filas con formato de texto: sin él, Excel convierte los códigos largos de SAP
# a número o a notación científica y se pierden dígitos.
TEXT_ROWS = 1000

# (encabezado, obligatorio, ancho, indicación)
COLUMNS = [
    ("Equipo", True, 14, "Código de equipo SAP (ej. 1020380). Si ese código ya existe en VALERA, la fila se omite."),
    ("Equipo superior", False, 16, "Código SAP del equipo padre. Debe existir en VALERA o venir en este mismo archivo."),
    ("Denominación", False, 36, "Descripción del equipo como aparece en SAP."),
    ("T", False, 6, "Tipo SAP: I = informática (dispositivo IT); H, M o V = equipo de operación. Vacío = equipo de operación."),
    ("Fabricante", False, 22, "Fabricante o marca."),
    ("Fabr. Nº-serie", False, 18, "Serial del fabricante."),
    ("NºPieza fabric.", False, 18, "Número de parte del fabricante."),
    ("Nº inventario", False, 16, "Placa física. Vacío = el N° de inventario en VALERA queda SAP-<Equipo>."),
    ("Activo fijo", False, 14, "Número de activo fijo SAP."),
    ("Sociedad", False, 10, "Sociedad SAP (ej. 1000)."),
    ("Centro coste", False, 13, "Centro de coste SAP."),
    ("Elemento PEP", False, 22, "Elemento PEP SAP."),
    ("Local", False, 16, "Base, pozo o rig."),
    ("Ubicac.técnica", False, 24, "Ubicación técnica SAP (ej. EV-PER-CEMTA-BATMIX)."),
    ("Emplazamiento", False, 14, "Emplazamiento SAP."),
    ("Ce.emplazam.", False, 13, "Centro del emplazamiento SAP."),
    ("Material", False, 14, "Código de material SAP."),
    ("Texto breve de material", False, 30, "Descripción del material SAP."),
    ("Status sistema", False, 14, "Status de sistema SAP (ej. AEQS)."),
]

INSTRUCTIONS = [
    "Esta plantilla da de alta equipos que todavía no existen en VALERA. Cada fila crea un activo en estado \"Por ubicar\".",
    "La columna Equipo (código SAP) es obligatoria y no puede repetirse. Si el código ya existe en VALERA, la fila se omite: "
    "para modificar activos existentes usa \"Exportar Excel\" y luego \"Editar en lote\".",
    "Equipo superior lleva el código SAP del equipo padre, que debe existir en VALERA o venir en este mismo archivo. "
    "Si no se encuentra, el activo se crea igual, sin padre, y el resultado del cargue lo indica.",
    "T define la línea del activo: I = dispositivo IT; H, M o V = equipo de operación. Vacío = equipo de operación.",
    "Cada encabezado tiene un comentario con lo que va en esa columna. Las columnas opcionales se pueden dejar vacías.",
    f"No cambies los encabezados ni el nombre de la hoja \"{ASSET_IMPORT_TEMPLATE_SHEET}\": VALERA los usa para reconocer el archivo.",
]


class AssetImportTemplateService(BaseExportService):

    def generate_file(self, data=None, options=None) -> io.BytesIO:
        wb = Workbook()
        self._build_data_sheet(wb.active)
        self._build_instructions_sheet(wb.create_sheet("Instrucciones"))

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _build_data_sheet(self, ws):
        ws.title = ASSET_IMPORT_TEMPLATE_SHEET
        ws.freeze_panes = "B2"

        for col, (label, required, width, hint) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = Font(bold=True, size=10, color=HEADER_FG)
            cell.fill = PatternFill(start_color=REQUIRED_BG if required else HEADER_BG, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
            cell.comment = Comment(("Obligatorio. " if required else "") + hint, "VALERA", width=260, height=110)
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = width
            for row in range(2, TEXT_ROWS + 2):
                ws.cell(row=row, column=col).number_format = "@"
        ws.row_dimensions[1].height = 30

        type_col = get_column_letter(next(i for i, c in enumerate(COLUMNS, start=1) if c[0] == "T"))
        type_validation = DataValidation(
            type="list",
            formula1='"I,H,M,V"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Tipo no válido",
            error="Usa I (informática) o H, M, V (equipo de operación), o deja la celda vacía.",
        )
        ws.add_data_validation(type_validation)
        type_validation.add(f"{type_col}2:{type_col}{TEXT_ROWS + 1}")

    def _build_instructions_sheet(self, ws):
        ws.column_dimensions["A"].width = 120
        title = ws.cell(row=1, column=1, value="Cómo llenar la plantilla de alta de activos")
        title.font = Font(bold=True, size=13, color=HEADER_FG)
        title.fill = PatternFill(start_color=HEADER_BG, fill_type="solid")
        ws.row_dimensions[1].height = 22

        for i, text in enumerate(INSTRUCTIONS, start=3):
            cell = ws.cell(row=i, column=1, value=f"{i - 2}. {text}")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[i].height = 32

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
