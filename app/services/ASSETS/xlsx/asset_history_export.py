import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ...base import BaseExportService
from ....models.AssetModel import AssetHistoryExportRequest

# Historial de UN activo como bitácora legible: cabecera con el activo como
# está hoy y una fila por movimiento, del más antiguo al más reciente. Los
# textos llegan ya resueltos en español desde valera; acá solo se diagrama.

NAVY = "1F3864"
WHITE = "FFFFFF"
LABEL_GRAY = "5B6779"
ZEBRA = "F7F9FC"

thin = Side(style="thin", color="D5DBE3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

COLUMNS = [
    ("Fecha", 17),
    ("Movimiento", 22),
    ("Detalle", 48),
    ("Estado anterior", 17),
    ("Estado nuevo", 17),
    ("Asignado a / Ubicación", 32),
    ("Condición / Nota", 38),
    ("Responsable", 26),
    ("Origen", 17),
]
LAST_COL = get_column_letter(len(COLUMNS))

# Color suave por tipo de movimiento: permite leer la bitácora de un vistazo
# (entregas, devoluciones y cambios de estado resaltan sobre las ediciones).
MOVEMENT_FILL = {
    "CREATED": "E6F4EA",
    "ASSIGNED": "FDEBDD",
    "RETURNED": "DDF3F7",
    "STATUS_CHANGED": "EFE6FB",
}
DEFAULT_MOVEMENT_FILL = "EEF2F7"

TITLE_ROW = 1
SUBTITLE_ROW = 2
INFO_FIRST_ROW = 4
META_ROW = 9
HEADER_ROW = 11
FIRST_DATA_ROW = HEADER_ROW + 1

DATA_FONT = Font(size=10)
DATA_BOLD = Font(size=10, bold=True)
ALIGN_TOP = Alignment(vertical="top", wrap_text=True)
ALIGN_TOP_CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _text(ws, row: int, column: int, value):
    """
    Celda de texto. openpyxl guarda como fórmula cualquier string que empiece
    con "=", y las condiciones y notas las escriben usuarios: forzar el tipo
    texto evita que un valor capturado se ejecute como fórmula al abrir el archivo.
    """
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


class AssetHistoryExportService(BaseExportService):

    def generate_file(self, data: AssetHistoryExportRequest, options=None) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Historial"
        ws.sheet_view.showGridLines = False

        for col, (_, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

        self._title(ws, data)
        self._asset_block(ws, data)
        self._meta(ws, data)
        self._table(ws, data)
        self._print_setup(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _title(self, ws, data: AssetHistoryExportRequest):
        ws.merge_cells(f"A{TITLE_ROW}:{LAST_COL}{TITLE_ROW}")
        cell = ws.cell(row=TITLE_ROW, column=1, value=f"HISTORIAL DEL ACTIVO {data.asset.assetTag}")
        cell.font = Font(bold=True, size=14, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[TITLE_ROW].height = 26

        ws.merge_cells(f"A{SUBTITLE_ROW}:{LAST_COL}{SUBTITLE_ROW}")
        sub = _text(ws, SUBTITLE_ROW, 1, data.asset.description or "Sin denominación")
        sub.font = Font(size=11, italic=True, color=LABEL_GRAY)
        sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[SUBTITLE_ROW].height = 18

    def _asset_block(self, ws, data: AssetHistoryExportRequest):
        a = data.asset
        brand_model = " · ".join(v for v in (a.brand, a.model) if v) or None
        # Dos columnas de pares etiqueta/valor: izquierda en A:C, derecha en E:I.
        left = [
            ("Nº inventario", a.assetTag),
            ("Equipo SAP", a.sapEquipmentCode),
            ("Tipo", a.assetLine),
            ("Fabricante / modelo", brand_model),
        ]
        right = [
            ("Estado actual", a.status),
            ("Asignado hoy a", a.currentAssignment or "Sin asignar"),
            ("Serial", a.serialNumber),
            ("Movimientos", str(len(data.rows))),
        ]
        for i, ((l_label, l_value), (r_label, r_value)) in enumerate(zip(left, right)):
            row = INFO_FIRST_ROW + i
            self._pair(ws, row, label_col=1, value_from=2, value_to=3, label=l_label, value=l_value)
            self._pair(ws, row, label_col=5, value_from=6, value_to=9, label=r_label, value=r_value)
            ws.row_dimensions[row].height = 17

    def _pair(self, ws, row: int, label_col: int, value_from: int, value_to: int, label: str, value):
        label_cell = ws.cell(row=row, column=label_col, value=label)
        label_cell.font = Font(size=10, bold=True, color=LABEL_GRAY)
        label_cell.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=row, start_column=value_from, end_row=row, end_column=value_to)
        value_cell = _text(ws, row, value_from, value or "—")
        value_cell.font = Font(size=10, bold=True)
        value_cell.alignment = Alignment(vertical="center")

    def _meta(self, ws, data: AssetHistoryExportRequest):
        ws.merge_cells(f"A{META_ROW}:{LAST_COL}{META_ROW}")
        by = f" por {data.generatedBy}" if data.generatedBy else ""
        cell = ws.cell(
            row=META_ROW,
            column=1,
            value=f"Generado el {data.generatedAt}{by}. Movimientos del más antiguo al más reciente.",
        )
        cell.font = Font(size=9, italic=True, color=LABEL_GRAY)

    def _table(self, ws, data: AssetHistoryExportRequest):
        for col, (label, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=HEADER_ROW, column=col, value=label)
            cell.font = Font(bold=True, size=10, color=WHITE)
            cell.fill = _fill(NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[HEADER_ROW].height = 22
        ws.freeze_panes = f"A{FIRST_DATA_ROW}"

        if not data.rows:
            ws.merge_cells(f"A{FIRST_DATA_ROW}:{LAST_COL}{FIRST_DATA_ROW}")
            empty = ws.cell(row=FIRST_DATA_ROW, column=1, value="Este activo no tiene movimientos registrados.")
            empty.font = Font(size=10, italic=True, color=LABEL_GRAY)
            empty.alignment = Alignment(horizontal="center")
            return

        zebra = _fill(ZEBRA)
        for i, row in enumerate(data.rows):
            r = FIRST_DATA_ROW + i
            values = [
                row.date,
                row.movement,
                row.detail,
                row.statusBefore or "",
                row.statusAfter or "",
                row.assignee or "",
                row.note or "",
                row.actor or "",
                row.source,
            ]
            for col, value in enumerate(values, start=1):
                cell = _text(ws, r, col, value)
                cell.border = BORDER
                cell.font = DATA_FONT
                cell.alignment = ALIGN_TOP_CENTER if col in (1, 4, 5, 9) else ALIGN_TOP
                if i % 2 == 1:
                    cell.fill = zebra

            movement = ws.cell(row=r, column=2)
            movement.font = DATA_BOLD
            movement.fill = _fill(MOVEMENT_FILL.get(row.type, DEFAULT_MOVEMENT_FILL))
            ws.cell(row=r, column=5).font = DATA_BOLD

        last_row = FIRST_DATA_ROW + len(data.rows) - 1
        ws.auto_filter.ref = f"A{HEADER_ROW}:{LAST_COL}{last_row}"

    def _print_setup(self, ws):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return ".xlsx"
