import io
import openpyxl
from typing import List, Dict, Any, Optional

# Formato estándar (2026-08-25): RH no logró producir de forma confiable el
# consolidado multi-hoja/multi-columna original de Nova — se adoptó un
# formato simple de una sola hoja con 2 columnas: documento y saldo. Nombre y
# empresa YA NO vienen del archivo — se resuelven del lado de VALERA
# (User/UserInformation/EmployeeContract) al momento de cargar, evitando
# depender de nombres de empresa inconsistentes entre Nova y VALERA.
DOC_HEADERS = {"documento", "codigoemp", "cedula", "cédula", "cc"}
ACCRUED_HEADERS = {"saldo", "cantidad", "vacaciones", "dias", "días"}


def _norm(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _norm_header(val) -> str:
    return _norm(val).lower()


def _find_col(headers: List[str], candidates: set) -> Optional[int]:
    for i, h in enumerate(headers):
        if h in candidates:
            return i
    return None


def _norm_document_id(val) -> str:
    """openpyxl entrega cédulas numéricas como int/float — normaliza sin '.0' final."""
    s = _norm(val)
    if not s:
        return ""
    try:
        return str(int(float(s)))
    except (ValueError, TypeError):
        return s


def _find_balance_sheet(wb):
    """
    Recorre las hojas buscando la que tenga columna de documento Y de saldo
    con datos reales — por encabezado, no por nombre/posición de hoja, para
    tolerar variaciones menores entre exports mensuales.
    """
    for ws in wb.worksheets:
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            continue
        headers = [_norm_header(c) for c in first_row]
        doc_idx = _find_col(headers, DOC_HEADERS)
        acc_idx = _find_col(headers, ACCRUED_HEADERS)
        if doc_idx is None or acc_idx is None:
            continue

        has_data = any(
            len(row) > acc_idx and _norm(row[acc_idx]) != ""
            for row in ws.iter_rows(min_row=2, values_only=True)
        )
        if not has_data:
            continue

        return ws, doc_idx, acc_idx
    return None


def parse_balance_import_file(file_bytes: bytes) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)

    found = _find_balance_sheet(wb)
    if found is None:
        raise ValueError(
            "No se encontró una hoja con columnas de documento y saldo "
            "(ej. documento/cedula y saldo/cantidad)"
        )
    ws, doc_idx, acc_idx = found

    rows_out: List[Dict[str, Any]] = []
    for raw_row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if all(v is None or _norm(v) == "" for v in row):
            continue

        row_number = raw_row_idx + 2
        errors: List[str] = []

        document_id = _norm_document_id(row[doc_idx] if doc_idx < len(row) else None)
        accrued_raw = row[acc_idx] if acc_idx < len(row) else None

        if not document_id:
            errors.append("Documento vacío")

        accrued: Optional[float] = None
        if accrued_raw is None or _norm(accrued_raw) == "":
            errors.append("Sin valor de saldo — fila omitida (ver hoja fuente)")
        else:
            try:
                accrued = float(accrued_raw)
            except (ValueError, TypeError):
                errors.append(f"Saldo inválido: '{accrued_raw}'")

        rows_out.append({
            "rowIndex": row_number,
            "documentId": document_id,
            "accrued": accrued,
            "parseErrors": errors,
        })

    return rows_out
